"""Minimal Markdown -> Word converter for the project's own documents.

Supports the subset the plan documents use: headings, paragraphs, bullet/numbered lists,
tables, fenced code blocks, blockquotes, bold and inline code. Chinese text gets an
East-Asian font so Word does not fall back to a random face.

    python tools/md_to_docx.py 计划书.md 计划书.docx

Requires python-docx (pip install python-docx); it is a documentation tool, not a
runtime dependency of the agent.
"""

from __future__ import annotations

import re
import sys
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor

BODY_FONT = "微软雅黑"
CODE_FONT = "Consolas"
BODY_SIZE = Pt(10.5)


def style_run(run, font: str = BODY_FONT, size: Pt = BODY_SIZE, bold: bool = False) -> None:
    run.font.name = font
    run.font.size = size
    run.font.bold = bold
    run._element.rPr.rFonts.set(qn("w:eastAsia"), font)


INLINE = re.compile(r"(\*\*.+?\*\*|`[^`]+`)")
CODE_SPAN = re.compile(r"(`[^`]+`)")


def add_inline(paragraph, text: str) -> None:
    """Render **bold** and `code` spans inside a paragraph."""
    for chunk in INLINE.split(text):
        if not chunk:
            continue
        if chunk.startswith("**") and chunk.endswith("**"):
            # Bold text may still contain `code` spans; handle one level of nesting.
            for part in CODE_SPAN.split(chunk[2:-2]):
                if not part:
                    continue
                if part.startswith("`") and part.endswith("`"):
                    run = paragraph.add_run(part[1:-1])
                    style_run(run, font=CODE_FONT, size=Pt(9.5), bold=True)
                    run.font.color.rgb = RGBColor(0xC0, 0x39, 0x2B)
                else:
                    style_run(paragraph.add_run(part), bold=True)
        elif chunk.startswith("`") and chunk.endswith("`"):
            run = paragraph.add_run(chunk[1:-1])
            style_run(run, font=CODE_FONT, size=Pt(9.5))
            run.font.color.rgb = RGBColor(0xC0, 0x39, 0x2B)
        else:
            style_run(paragraph.add_run(chunk))


def add_code_block(doc: Document, lines: list[str]) -> None:
    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.space_before = Pt(4)
    paragraph.paragraph_format.space_after = Pt(8)
    paragraph.paragraph_format.left_indent = Pt(12)
    for index, line in enumerate(lines):
        run = paragraph.add_run(line)
        style_run(run, font=CODE_FONT, size=Pt(9))
        if index != len(lines) - 1:
            run.add_break()


def add_table(doc: Document, rows: list[list[str]]) -> None:
    if not rows:
        return
    width = max(len(row) for row in rows)
    table = doc.add_table(rows=0, cols=width)
    table.style = "Table Grid"
    for row_index, row in enumerate(rows):
        cells = table.add_row().cells
        for column in range(width):
            text = row[column] if column < len(row) else ""
            paragraph = cells[column].paragraphs[0]
            add_inline(paragraph, text)
            for run in paragraph.runs:
                run.font.size = Pt(9.5)
                if row_index == 0:
                    run.font.bold = True
    doc.add_paragraph()


def split_row(line: str) -> list[str]:
    return [cell.strip() for cell in line.strip().strip("|").split("|")]


def is_separator(line: str) -> bool:
    return bool(re.fullmatch(r"\|[\s:\-|]+\|", line.strip()))


def convert(markdown: str, doc: Document) -> None:
    lines = markdown.splitlines()
    index = 0
    while index < len(lines):
        line = lines[index]
        stripped = line.strip()

        if stripped.startswith("```"):
            index += 1
            block: list[str] = []
            while index < len(lines) and not lines[index].strip().startswith("```"):
                block.append(lines[index])
                index += 1
            add_code_block(doc, block)
            index += 1
            continue

        if stripped.startswith("|") and index + 1 < len(lines) and is_separator(lines[index + 1]):
            rows: list[list[str]] = []
            while index < len(lines) and lines[index].strip().startswith("|"):
                if not is_separator(lines[index]):
                    rows.append(split_row(lines[index]))
                index += 1
            add_table(doc, rows)
            continue

        if stripped.startswith("#### ") or stripped.startswith("### "):
            heading = doc.add_heading(level=3)
            add_inline(heading, stripped.lstrip("#").strip())
            index += 1
            continue
        if stripped.startswith("## "):
            heading = doc.add_heading(level=2)
            add_inline(heading, stripped[3:].strip())
            index += 1
            continue
        if stripped.startswith("# "):
            heading = doc.add_heading(level=1)
            add_inline(heading, stripped[2:].strip())
            index += 1
            continue

        if stripped.startswith("> "):
            paragraph = doc.add_paragraph()
            paragraph.paragraph_format.left_indent = Pt(18)
            add_inline(paragraph, stripped[2:])
            for run in paragraph.runs:
                run.font.italic = True
            index += 1
            continue

        if re.match(r"^[-*] ", stripped):
            paragraph = doc.add_paragraph(style="List Bullet")
            add_inline(paragraph, stripped[2:])
            index += 1
            continue

        if re.match(r"^\d+\. ", stripped):
            paragraph = doc.add_paragraph(style="List Number")
            add_inline(paragraph, re.sub(r"^\d+\. ", "", stripped))
            index += 1
            continue

        if stripped in ("---", "***", ""):
            index += 1
            continue

        paragraph = doc.add_paragraph()
        paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
        add_inline(paragraph, stripped)
        index += 1


def main(argv: list[str]) -> int:
    if len(argv) < 2:
        print(__doc__)
        return 2
    source = Path(argv[1])
    target = Path(argv[2]) if len(argv) > 2 else source.with_suffix(".docx")
    doc = Document()
    convert(source.read_text(encoding="utf-8"), doc)
    doc.save(str(target))
    print(f"wrote {target}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main(sys.argv))
